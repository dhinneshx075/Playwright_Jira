"""Outputs of a run: results JSON (same shape as the input) and a Word evidence document."""

from __future__ import annotations

import copy
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, RGBColor

from .models import TestCase

STATUS_COLORS = {
    "PASS": RGBColor(0x1E, 0x7B, 0x34),
    "FAIL": RGBColor(0xC0, 0x1C, 0x28),
    "WARN": RGBColor(0xB2, 0x6B, 0x00),
    "BLOCKED": RGBColor(0x6B, 0x6B, 0x6B),
    "NOT_AUTOMATED": RGBColor(0x6B, 0x6B, 0x6B),
    "NO_DATA": RGBColor(0x9A, 0x5B, 0x00),
    "SKIPPED": RGBColor(0x6B, 0x6B, 0x6B),
    "PARTIAL": RGBColor(0xB2, 0x6B, 0x00),
}


def write_results_json(test_case: TestCase, run_dir: Path) -> Path:
    """The input JSON with actual_result/status/screenshots filled in for every step."""
    data = copy.deepcopy(test_case.raw)

    data["automation_result"] = {
        "status": test_case.status,
        "started_at": test_case.started_at,
        "finished_at": test_case.finished_at,
        "evidence_folder": str(run_dir),
        "downloads": [str(p) for p in test_case.downloads],
    }

    for raw_step, step in zip(data.get("steps", []), test_case.steps):
        raw_step["actual_result"] = step.actual_result
        raw_step["status"] = step.status
        raw_step["screenshots"] = [str(p) for p in step.screenshots]
        raw_step["action_results"] = [
            {
                "index": r.index,
                "action": r.action,
                "description": r.description,
                "status": r.status,
                "message": r.message,
                "screenshot": str(r.screenshot) if r.screenshot else None,
            }
            for r in step.action_results
        ]

    path = run_dir / f"result_{test_case.id}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    return path


def _status_paragraph(document, label: str, status: str):
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    run = paragraph.add_run(status)
    run.bold = True
    if status in STATUS_COLORS:
        run.font.color.rgb = STATUS_COLORS[status]


def _labeled(document, label: str, text: str):
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(text or "-")


def create_evidence_document(test_case: TestCase, run_dir: Path) -> Path:
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = document.add_heading("Test Execution Evidence", level=0)
    title.alignment = 1

    summary = document.add_table(rows=0, cols=2)
    summary.style = "Light Grid Accent 1"
    for label, value in [
        ("Test Cycle", test_case.cycle_key),
        ("Assigned Test Case ID", test_case.assigned_test_case_id),
        ("Project ID", str(test_case.raw.get("project_id", ""))),
        ("Jira URL", test_case.raw.get("url", "")),
        ("Application URL", test_case.app_url),
        ("Started", test_case.started_at),
        ("Finished", test_case.finished_at),
        ("Overall Status", test_case.status),
        ("Downloaded Files", ", ".join(p.name for p in test_case.downloads)),
    ]:
        cells = summary.add_row().cells
        cells[0].text = label
        cells[1].text = value or ""

    for step in test_case.steps:
        document.add_page_break()
        jira_step = step.raw.get("jira_step")
        document.add_heading(f"Step {step.sequence}"
                             + (f" (Jira step {jira_step})" if jira_step else ""), level=1)

        _status_paragraph(document, "Status", step.status)
        _labeled(document, "Step", step.step)
        _labeled(document, "Test Data", step.test_data)
        _labeled(document, "Expected Result", step.expected_result)
        _labeled(document, "Actual Result", step.actual_result.split("\n")[0])

        if step.action_results:
            document.add_heading("Automated Checks", level=2)
            table = document.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            header = table.rows[0].cells
            for cell, text in zip(header, ["#", "Action", "Status", "Details"]):
                cell.text = text
            for result in step.action_results:
                cells = table.add_row().cells
                cells[0].text = str(result.index)
                cells[1].text = result.action + (
                    f"\n{result.description}" if result.description else ""
                )
                cells[2].text = result.status
                cells[3].text = result.message

        if step.screenshots:
            document.add_heading("Screenshots", level=2)
            for path in step.screenshots:
                if Path(path).exists():
                    document.add_paragraph().add_run(Path(path).name).italic = True
                    document.add_picture(str(path), width=Inches(6.8))

    path = run_dir / f"Evidence_{test_case.id}.docx"
    document.save(str(path))
    return path

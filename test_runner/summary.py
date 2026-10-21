"""
Run summary: one Word document (plus a small JSON) for everything a run command did -
every test case run, and every one that could not be run - written to test_runs\\.
"""

from __future__ import annotations

import getpass
import json
import os
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .models import TestCase
from .report import STATUS_COLORS

NOT_RUN = "NOT_RUN"
NO_DATA = "NO_DATA"
_FILL = {"PASS": "E3F4E8", "FAIL": "FBE3E5", "PARTIAL": "FFF1DC", "NOT_AUTOMATED": "EEEEEE",
         NOT_RUN: "EEEEEE", NO_DATA: "FFF1DC"}
_ACCENT = RGBColor(0x1F, 0x4E, 0x79)
_GREY = RGBColor(0x55, 0x55, 0x55)


@dataclass
class SummaryEntry:
    """One test case in the run: run (test_case set) or not run (reason set)."""
    label: str
    status: str
    test_case: TestCase | None = None
    evidence: Path | None = None
    reason: str = ""


@dataclass
class RunSummary:
    command: str
    started: datetime = field(default_factory=datetime.now)
    finished: datetime | None = None
    entries: list[SummaryEntry] = field(default_factory=list)

    def add_run(self, test_case: TestCase, evidence: Path):
        self.entries.append(SummaryEntry(test_case.id, test_case.status, test_case, evidence))

    def add_not_run(self, label: str, reason: str):
        self.entries.append(SummaryEntry(label, NOT_RUN, reason=reason))


# ---------------------------------------------------------------------------- facts

def _name(entry: SummaryEntry) -> str:
    return str(entry.test_case.raw.get("name", "")) if entry.test_case else ""


def _duration(test_case: TestCase | None) -> str:
    try:
        seconds = int((datetime.fromisoformat(test_case.finished_at)
                       - datetime.fromisoformat(test_case.started_at)).total_seconds())
    except (ValueError, TypeError, AttributeError):
        return "-"
    return f"{seconds // 60}m {seconds % 60:02d}s" if seconds >= 60 else f"{seconds}s"


def _steps(test_case: TestCase | None) -> str:
    if not test_case:
        return "-"
    passed = sum(1 for s in test_case.steps if s.status == "PASS")
    return f"{passed}/{len(test_case.steps)}"


def _failed_step(test_case: TestCase):
    for step in test_case.steps:
        if step.status == "FAIL":
            return step
    return next((s for s in test_case.steps if s.status == "BLOCKED"), None)


def _result_text(entry: SummaryEntry) -> str:
    if entry.status == NOT_RUN:
        return entry.reason
    test_case = entry.test_case
    if entry.status == NO_DATA:
        return f"Report returned no data: {test_case.no_data_message or 'No data found'}"
    if entry.status == "PASS":
        return "All automated checks passed."
    if entry.status in ("PARTIAL", "NOT_AUTOMATED"):
        manual = sum(1 for s in test_case.steps if s.status == "NOT_AUTOMATED")
        return f"{manual} step(s) must be executed manually."
    step = _failed_step(test_case)
    if step is None:
        return ""
    first_line = (step.actual_result or "").split("\n")[0]
    return f"Step {step.sequence}: {first_line}"


def _failure_screenshot(test_case: TestCase) -> Path | None:
    step = _failed_step(test_case)
    if step is None:
        return None
    for result in step.action_results:
        if result.status == "FAIL" and result.screenshot and Path(result.screenshot).exists():
            return Path(result.screenshot)
    shots = [p for p in step.screenshots if Path(p).exists()]
    return Path(shots[-1]) if shots else None


def _counts(entries):
    count = {s: sum(1 for e in entries if e.status == s)
             for s in ("PASS", "FAIL", "PARTIAL", "NOT_AUTOMATED", NO_DATA, NOT_RUN)}
    count["total"] = len(entries)
    count["testable"] = count["total"] - count[NOT_RUN] - count[NO_DATA] - count["NOT_AUTOMATED"]
    count["rate"] = (f"{round(100 * count['PASS'] / count['testable'])}%"
                      if count["testable"] else "-")
    return count


def _overall(count) -> str:
    if count["FAIL"] or count[NOT_RUN]:
        return "FAIL"
    if count["PARTIAL"] or count["NOT_AUTOMATED"] or count[NO_DATA]:
        return "PARTIAL"
    return "PASS"


# ---------------------------------------------------------------------------- docx helpers

def _shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_text(cell, text: str, bold=False, color=None, size=9, align=None):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return paragraph


def _hyperlink(paragraph, text: str, target: Path, size=9):
    """A clickable link to a local file."""
    r_id = paragraph.part.relate_to(target.resolve().as_uri(), RELATIONSHIP_TYPE.HYPERLINK,
                                    is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, value in (("w:color", "1F4E79"), ("w:u", "single"), ("w:sz", str(size * 2))):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        props.append(element)
    run.append(props)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


def _table(document, headers, widths_cm):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, header in zip(table.rows[0].cells, headers):
        _shade(cell, "1F4E79")
        _cell_text(cell, header, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    table._widths = widths_cm
    return table


def _fix_widths(table):
    for row in table.rows:
        for cell, width in zip(row.cells, table._widths):
            cell.width = Cm(width)


def _heading(document, text: str, level: int):
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = _ACCENT
    return heading


# ---------------------------------------------------------------------------- document

def create_run_summary(run: RunSummary, output_root: Path) -> tuple[Path, Path]:
    """Write Run_Summary_<time>.docx and .json into output_root. Returns both paths."""
    run.finished = run.finished or datetime.now()
    output_root.mkdir(parents=True, exist_ok=True)
    stamp = run.started.strftime("%Y%m%d_%H%M%S")
    count = _counts(run.entries)
    overall = _overall(count)

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(1.5))
    section.top_margin = section.bottom_margin = Cm(1.3)

    # Title
    title = document.add_paragraph()
    run_title = title.add_run("Test Execution Summary")
    run_title.bold = True
    run_title.font.size = Pt(24)
    run_title.font.color.rgb = _ACCENT
    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(
        f"{run.started.strftime('%d %B %Y')}   |   Overall result: ")
    subtitle_run.font.color.rgb = _GREY
    verdict = subtitle.add_run(overall)
    verdict.bold = True
    verdict.font.color.rgb = STATUS_COLORS.get(overall, _GREY)

    # Run details
    sites = sorted({e.test_case.app_url for e in run.entries if e.test_case})
    duration = int((run.finished - run.started).total_seconds())
    details = _table(document, ["Run details", ""], [5.0, 21.7])
    for label, value in [
        ("Scope", run.command),
        ("Application tested", ", ".join(sites) or "-"),
        ("Executed by", f"{getpass.getuser()} on {os.environ.get('COMPUTERNAME', '-')}"),
        ("Started / finished", f"{run.started:%d %b %Y %H:%M:%S}  -  {run.finished:%H:%M:%S}"),
        ("Duration", f"{duration // 60}m {duration % 60:02d}s"),
        ("Evidence folder", str(output_root.resolve())),
    ]:
        cells = details.add_row().cells
        _cell_text(cells[0], label, bold=True)
        _cell_text(cells[1], value)
    merged = details.rows[0].cells[0].merge(details.rows[0].cells[1])
    for extra in merged.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _fix_widths(details)
    document.add_paragraph()

    # Result tiles
    tiles = [("Total", count["total"], "1F4E79", None),
             ("Passed", count["PASS"], _FILL["PASS"], "PASS"),
             ("Failed", count["FAIL"], _FILL["FAIL"], "FAIL"),
             ("Partial / manual", count["PARTIAL"] + count["NOT_AUTOMATED"], _FILL["PARTIAL"], "PARTIAL"),
             ("No report data", count[NO_DATA], _FILL[NO_DATA], NO_DATA),
             ("Not run", count[NOT_RUN], _FILL[NOT_RUN], NOT_RUN),
             ("Pass rate (of executed)", count["rate"], "1F4E79", None)]
    tile_table = document.add_table(rows=1, cols=len(tiles))
    tile_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (label, value, fill, status) in zip(tile_table.rows[0].cells, tiles):
        _shade(cell, fill)
        white = RGBColor(0xFF, 0xFF, 0xFF) if status is None else None
        number = _cell_text(cell, str(value), bold=True, size=22,
                            color=white or STATUS_COLORS.get(status, _GREY),
                            align=WD_ALIGN_PARAGRAPH.CENTER)
        number.paragraph_format.space_before = Pt(4)
        caption = cell.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(4)
        caption_run = caption.add_run(label)
        caption_run.font.size = Pt(9)
        caption_run.font.color.rgb = white or _GREY
        cell.width = Cm(26.7 / len(tiles))

    # Results table
    _heading(document, "Results", 1)
    results = _table(document, ["#", "Test case", "Name", "Status", "Steps passed", "Duration",
                                "Result / reason", "Evidence"],
                     [0.8, 4.4, 5.4, 2.4, 1.7, 1.7, 7.6, 2.7])
    for number, entry in enumerate(run.entries, start=1):
        cells = results.add_row().cells
        _cell_text(cells[0], str(number))
        _cell_text(cells[1], entry.label, bold=True)
        _cell_text(cells[2], _name(entry))
        _shade(cells[3], _FILL.get(entry.status, "FFFFFF"))
        _cell_text(cells[3], entry.status.replace("_", " "), bold=True,
                   color=STATUS_COLORS.get(entry.status, _GREY))
        _cell_text(cells[4], _steps(entry.test_case))
        _cell_text(cells[5], _duration(entry.test_case))
        _cell_text(cells[6], _result_text(entry))
        paragraph = cells[7].paragraphs[0]
        if entry.evidence and entry.evidence.exists():
            _hyperlink(paragraph, "Open evidence", entry.evidence)
        else:
            _cell_text(cells[7], "-")
    _fix_widths(results)

    # Failures in detail
    failed = [e for e in run.entries if e.status == "FAIL" and e.test_case]
    if failed:
        document.add_page_break()
        _heading(document, "Failures", 1)
        for entry in failed:
            _heading(document, f"{entry.label}  {_name(entry)}".strip(), 2)
            step = _failed_step(entry.test_case)
            if step is not None:
                failing = next((r for r in step.action_results if r.status == "FAIL"), None)
                for label, text in [
                    ("Failed at", f"Step {step.sequence}"
                     + (f", command {failing.index}: {failing.description or failing.action}"
                        if failing else "")),
                    ("Message", failing.message if failing else step.actual_result.split("\n")[0]),
                    ("Expected", step.expected_result or "-"),
                ]:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.add_run(f"{label}: ").bold = True
                    paragraph.add_run(text)
            screenshot = _failure_screenshot(entry.test_case)
            if screenshot is not None:
                document.add_picture(str(screenshot), width=Inches(7.5))
            if entry.evidence and entry.evidence.exists():
                paragraph = document.add_paragraph()
                paragraph.add_run("Full evidence: ")
                _hyperlink(paragraph, entry.evidence.name, entry.evidence, size=10)

    no_data = [e for e in run.entries if e.status == NO_DATA]
    if no_data:
        _heading(document, "Reports with no data", 1)
        table = _table(document, ["Test case", "Result", "Evidence"], [4.5, 18.2, 2.7])
        for entry in no_data:
            cells = table.add_row().cells
            _cell_text(cells[0], entry.label, bold=True)
            _cell_text(cells[1], _result_text(entry))
            if entry.evidence and entry.evidence.exists():
                _hyperlink(cells[2].paragraphs[0], "Open evidence", entry.evidence)
            else:
                _cell_text(cells[2], "-")
        _fix_widths(table)

    # Not executed
    not_run = [e for e in run.entries if e.status == NOT_RUN]
    if not_run:
        _heading(document, "Not executed", 1)
        table = _table(document, ["Test case / source", "Reason"], [6.0, 20.7])
        for entry in not_run:
            cells = table.add_row().cells
            _cell_text(cells[0], entry.label, bold=True)
            _cell_text(cells[1], entry.reason)
        _fix_widths(table)

    # Footer
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"Test Execution Summary  |  {run.started:%d %b %Y %H:%M}  |  "
                                "generated by test_runner")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = _GREY

    docx_path = output_root / f"Run_Summary_{stamp}.docx"
    document.save(str(docx_path))

    # Machine-readable copy (for run history / trends)
    json_path = output_root / f"Run_Summary_{stamp}.json"
    json_path.write_text(json.dumps({
        "scope": run.command,
        "started": run.started.isoformat(timespec="seconds"),
        "finished": run.finished.isoformat(timespec="seconds"),
        "overall": overall,
        "counts": {k: v for k, v in count.items()},
        "test_cases": [{
            "id": e.label,
            "name": _name(e),
            "status": e.status,
            "steps_passed": _steps(e.test_case),
            "duration": _duration(e.test_case),
            "result": _result_text(e),
            "evidence": str(e.evidence) if e.evidence else None,
        } for e in run.entries],
    }, indent=2, ensure_ascii=False), encoding="utf-8")
    return docx_path, json_path
